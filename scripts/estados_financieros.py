#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
estados_financieros.py
======================
Genera el paquete completo de estados financieros de una empresa ERPNext y lo
exporta a HTML / CSV / XLSX en un solo paso:

    01  Balance General                    (Estado de Situación Financiera)
    02  Estado de Resultados               (Pérdidas y Ganancias)
    03  Balance Financiero                 (Balanza de Comprobación)
    04  Estado de Flujo de Efectivo
    05  Estado de Cambios en el Patrimonio Neto  (formato matricial NIC 1)
    05b Estado de Cambios en el Patrimonio Neto  (plantilla nativa ERPNext)
    06  Libro Diario                       (formato Guatemala, partidas con folio)
    07  Libro Mayor                        (formato Guatemala, saldo corriente)

Uso (lo hace 'install.sh reports'):

    docker compose exec -T backend bench --site <sitio> console \
        < scripts/estados_financieros.py

Parámetros por variables de entorno:
    EF_SITE, EF_COMPANY, EF_FISCAL_YEAR, EF_OUTDIR, LD_FOLIO_INICIO
"""

import csv
import html
import os
import sys
from datetime import date

# --------------------------------------------------------------------------
# Bootstrap
# --------------------------------------------------------------------------


def _bootstrap():
    import frappe

    already = getattr(frappe.local, "site", None) and getattr(frappe.local, "db", None)
    if already:
        return frappe

    site = os.environ.get("EF_SITE") or os.environ.get("ERP_SITE") or (
        sys.argv[1] if len(sys.argv) > 1 else None
    )
    sites_path = os.environ.get("ERP_SITES_PATH", "/home/frappe/frappe-bench/sites")
    frappe.init(site=site, sites_path=sites_path)
    frappe.connect()
    return frappe


frappe = _bootstrap()

from frappe.utils import add_days, cstr, flt, getdate  # noqa: E402

frappe.set_user("Administrator")
frappe.flags.ignore_permissions = True

SENTINEL = "ERP_FINANCIAL_STATEMENTS_OK"
OUTDIR = os.environ.get("EF_OUTDIR") or "/home/frappe/frappe-bench/sites/_estados_financieros"

# Rellenados en main(); usados por _emit() para el encabezado de cada archivo.
SUBTITLE = ""
IDENTITY = {"name": "", "tax_id": "", "currency": ""}


# --------------------------------------------------------------------------
# Contexto (empresa y ejercicio)
# --------------------------------------------------------------------------


def resolve_company():
    name = os.environ.get("EF_COMPANY")
    if name and frappe.db.exists("Company", name):
        return name
    if name:
        found = frappe.db.get_value("Company", {"company_name": name}, "name")
        if found:
            return found
    name = frappe.defaults.get_global_default("company")
    if name and frappe.db.exists("Company", name):
        return name
    companies = frappe.get_all("Company", pluck="name", limit=1)
    if not companies:
        raise SystemExit(
            "No hay ninguna empresa todavía.\n"
            "  1) Abre el navegador en el sitio y crea tu empresa con el asistente\n"
            "     (País: Guatemala, Moneda: GTQ, plan: Guatemala - Cuentas).\n"
            "  2) Ejecuta:  ./install.sh contabilidad\n"
            "  3) Vuelve a intentar:  ./install.sh reports"
        )
    return companies[0]


def resolve_fiscal_year(company, label=None):
    from erpnext.accounts.utils import get_fiscal_year

    fy = None
    if label:
        fy = get_fiscal_year(
            fiscal_year=label, company=company, as_dict=True, raise_on_missing=False
        )
    if not fy:
        fy = get_fiscal_year(
            date=date.today(), company=company, as_dict=True, raise_on_missing=False
        )
    if not fy:
        fys = frappe.get_all(
            "Fiscal Year",
            fields=["name", "year_start_date", "year_end_date"],
            order_by="year_start_date desc",
            limit=1,
        )
        if not fys:
            raise SystemExit("No hay años fiscales definidos. Ejecuta setup_accounting.py.")
        fy = frappe._dict(fys[0])
    return fy


def d2s(value):
    return getdate(value).strftime("%Y-%m-%d")


# --------------------------------------------------------------------------
# Ejecución de reportes
# --------------------------------------------------------------------------


def run_report(report_name, filters):
    from frappe.desk.query_report import run

    out = run(report_name, filters=dict(filters))
    if not isinstance(out, dict):
        return [], []
    return out.get("columns") or [], out.get("result") or []


def fin_filters(company, from_date, to_date, fy, **extra):
    filters = frappe._dict(
        {
            "company": company,
            "filter_based_on": "Date Range",
            "period_start_date": d2s(from_date),
            "period_end_date": d2s(to_date),
            "from_fiscal_year": fy,
            "to_fiscal_year": fy,
            "periodicity": "Yearly",
            "accumulated_values": 1,
            "include_default_book_entries": 1,
            "selected_view": "Report",
            "show_zero_values": 1,
            "show_account_details": "Summary",
        }
    )
    filters.update(extra)
    return filters


# --------------------------------------------------------------------------
# Exportación
# --------------------------------------------------------------------------


def fmt_value(value, column):
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, (int, float)):
        fieldtype = (column or {}).get("fieldtype")
        if fieldtype == "Int":
            return str(int(flt(value, 0)))
        if fieldtype == "Percent":
            return f"{flt(value, 2):.2f}%"
        return f"{flt(value, 2):,.2f}"
    return cstr(value)


def write_csv(path, columns, rows):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow([c.get("label") or c.get("fieldname") for c in columns])
        for row in rows:
            writer.writerow([fmt_value(row.get(c.get("fieldname")), c) for c in columns])


# --------------------------------------------------------------------------
# Presentación HTML (diseño compartido por los estados y los libros)
# --------------------------------------------------------------------------

_PAGE_CSS = """
:root{
  --ink:#1b2831; --muted:#6a7883; --line:#e3eaf0; --soft:#f6f9fb; --soft2:#eaf1f6;
  --brand:#0d5c63; --brand2:#149089; --gold:#b3861f; --neg:#b4302a; --pos:#177245;
}
*{box-sizing:border-box}
html,body{margin:0}
body{background:#e8eef3;color:var(--ink);
     font-family:-apple-system,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;
     -webkit-print-color-adjust:exact;print-color-adjust:exact}
.sheet{max-width:1120px;margin:26px auto;background:#fff;border:1px solid var(--line);
       border-radius:16px;overflow:hidden;box-shadow:0 18px 44px rgba(13,45,60,.10)}
.masthead{background:linear-gradient(135deg,var(--brand),var(--brand2));color:#fff;
          padding:22px 32px 20px;position:relative}
.masthead::after{content:"";position:absolute;left:0;right:0;bottom:0;height:3px;
                 background:linear-gradient(90deg,var(--gold),transparent 70%)}
.masthead .co{font-size:12px;letter-spacing:.16em;text-transform:uppercase;opacity:.88}
.masthead h1{margin:6px 0 3px;font-size:23px;letter-spacing:.02em;font-weight:650}
.masthead .period{font-size:12.5px;opacity:.92}
.content{padding:22px 32px 28px}
.cards{display:flex;flex-wrap:wrap;gap:12px;margin:0 0 20px}
.card{flex:1 1 168px;min-width:146px;border:1px solid var(--line);border-radius:12px;
      padding:12px 14px;background:linear-gradient(180deg,#fff,var(--soft))}
.card .k{display:block;font-size:10.5px;letter-spacing:.09em;text-transform:uppercase;color:var(--muted)}
.card .v{display:block;font-size:19px;font-weight:650;margin-top:5px;font-variant-numeric:tabular-nums}
.card.neg .v{color:var(--neg)} .card.pos .v{color:var(--pos)}
table{border-collapse:collapse;width:100%;font-size:13px}
thead th{background:var(--soft2);color:#33454f;text-align:left;font-weight:600;
         padding:9px 11px;border-bottom:1px solid var(--line);white-space:nowrap}
tbody td{padding:8px 11px;border-bottom:1px solid var(--line);vertical-align:top}
tbody tr:nth-child(even) td{background:#fbfdfe}
tbody tr:hover td{background:#f1f7fa}
td.num,th.num{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
td.cod{color:#4c5a64}
tr.total td{font-weight:700;background:#f1f8f5;border-top:2px solid #cfe4de}
tr.total td:first-child{text-transform:uppercase;letter-spacing:.02em;font-size:12px}
.neg{color:var(--neg)}
.note{margin-top:18px;color:var(--muted);font-size:11.5px;border-top:1px solid var(--line);padding-top:10px}
.badge{display:inline-block;padding:3px 11px;border-radius:999px;font-size:11.5px;
       font-weight:650;letter-spacing:.04em}
.badge.ok{background:#e3f4ea;color:var(--pos);border:1px solid #bfe3cd}
.badge.bad{background:#fbe6e4;color:var(--neg);border:1px solid #f0c6c2}
.panel{margin-top:22px;border:1px solid var(--line);border-radius:12px;overflow:hidden}
.panel > h3{margin:0;padding:10px 15px;background:var(--soft2);font-size:12px;letter-spacing:.08em;
            text-transform:uppercase;color:#33454f;border-bottom:1px solid var(--line)}
.panel .body{padding:14px 15px;border-top:1px solid var(--line)}
.panel .body p{margin:0}
.asiento{margin:0 0 16px;border:1px solid var(--line);border-radius:12px;overflow:hidden;
         page-break-inside:avoid}
.ahead{display:flex;flex-wrap:wrap;gap:5px 18px;align-items:baseline;background:var(--soft);
       border-bottom:1px solid var(--line);padding:9px 14px;font-size:12.5px;color:#3b4b55}
.ahead .folio{font-weight:700;color:var(--brand)}
.ahead .doc{color:var(--muted)}
.acc{margin:0 0 18px;page-break-inside:avoid}
.acc h2{font-size:13.5px;margin:0 0 6px;color:var(--brand);letter-spacing:.02em}
.acc h2 .code{color:var(--muted);font-weight:600;margin-right:6px}
.acc tfoot td{font-weight:650;background:var(--soft2);border-top:2px solid #cfe4de}
h2.sec{font-size:14px;margin:24px 0 10px;color:#33454f;letter-spacing:.02em}
.content ul{margin:0;padding-left:20px}
.content li{margin:7px 0}
.muted{color:var(--muted)}
a{color:var(--brand2)}
@media print{
  body{background:#fff}
  .sheet{margin:0;border:none;border-radius:0;box-shadow:none;max-width:none}
  thead{display:table-header-group}
  tr,.asiento,.acc{page-break-inside:avoid}
}
"""


def _esc(value):
    return html.escape(cstr(value))


def _money(value):
    return f"{flt(value):,.2f}"


def _num_cell(value):
    """Celda numérica; los saldos negativos se pintan en rojo (clase .neg)."""
    cls = "num neg" if flt(value) < 0 else "num"
    return f'<td class="{cls}">{_money(value)}</td>'


def _page(title, body, extra_css=""):
    return (
        '<!doctype html>\n<html lang="es"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f"<title>{_esc(title)}</title>"
        f"<style>{_PAGE_CSS}{extra_css}</style></head><body>"
        f'<div class="sheet">{body}</div></body></html>'
    )


def _masthead(title, period):
    name = IDENTITY.get("name") or ""
    tax = f" · NIT {_esc(IDENTITY['tax_id'])}" if IDENTITY.get("tax_id") else ""
    lead = f"{_esc(name)}{tax}" if name else "ERPNext · Contabilidad"
    period_html = f'<div class="period">{_esc(period)}</div>' if period else ""
    return (
        '<header class="masthead">'
        f'<div class="co">{lead}</div>'
        f"<h1>{_esc(title)}</h1>{period_html}"
        "</header>"
    )


def _cards_html(cards):
    items = []
    for label, value in cards or []:
        if value is None:
            continue
        if isinstance(value, str):
            text, cls = value, "card"
        else:
            number = flt(value)
            text = _money(number)
            cls = "card neg" if number < 0 else "card"
        items.append(
            f'<div class="{cls}"><span class="k">{_esc(label)}</span>'
            f'<span class="v">{_esc(text)}</span></div>'
        )
    return f'<div class="cards">{"".join(items)}</div>' if items else ""


def _badge(ok, ok_text="CUADRA", bad_text="NO CUADRA"):
    if ok:
        return f'<span class="badge ok">{_esc(ok_text)}</span>'
    return f'<span class="badge bad">{_esc(bad_text)}</span>'


_TOTAL_PREFIXES = ("total", "grand total", "suma", "utilidad", "pérdida", "perdida", "resultado")


def _is_total_row(label, row=None):
    """¿La fila es un total? (para resaltarla en negrita)."""
    if row and row.get("is_total_row"):
        return True
    low = cstr(label).strip().lower()
    if not low:
        return False
    return (
        low.startswith(_TOTAL_PREFIXES)
        or "profit for the year" in low
        or low.startswith("net ")
    )


def write_html(path, title, subtitle, columns, rows, note="", cards=None):
    """Escribe un reporte como HTML imprimible con el diseño compartido."""
    body = _masthead(title, subtitle) + '<div class="content">' + _cards_html(cards)

    if columns:
        head_cells = []
        for idx, col in enumerate(columns):
            label = _esc(col.get("label") or col.get("fieldname"))
            head_cells.append(f'<th class="num">{label}</th>' if idx else f"<th>{label}</th>")
        heads = "".join(head_cells)
        first_field = columns[0].get("fieldname")
        trs = []
        for row in rows or []:
            is_total = _is_total_row(row.get(first_field), row)
            cells = []
            for idx, col in enumerate(columns):
                value = row.get(col.get("fieldname"))
                text = fmt_value(value, col)
                cls = []
                if idx:
                    cls.append("num")
                if isinstance(value, (int, float)) and not isinstance(value, bool) and flt(value) < 0:
                    cls.append("neg")
                klass = f' class="{" ".join(cls)}"' if cls else ""
                cells.append(f"<td{klass}>{html.escape(text)}</td>")
            if is_total:
                trs.append(f'<tr class="total">{"".join(cells)}</tr>')
            else:
                trs.append(f'<tr>{"".join(cells)}</tr>')
        body += f'<table><thead><tr>{heads}</tr></thead><tbody>{"".join(trs)}</tbody></table>'

    if note:
        body += f'<div class="note">{_esc(note)}</div>'
    body += "</div>"

    with open(path, "w", encoding="utf-8") as handle:
        handle.write(_page(title, body))


def write_xlsx(path, sheets):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception:  # noqa: BLE001
        return False

    workbook = Workbook()
    workbook.remove(workbook.active)
    head_fill = PatternFill("solid", fgColor="F6F8FA")
    for name, columns, rows in sheets:
        sheet = workbook.create_sheet(title=name[:31])
        sheet.append([c.get("label") or c.get("fieldname") for c in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
        for row in rows:
            sheet.append([row.get(c.get("fieldname")) for c in columns])
        for idx, col in enumerate(columns, start=1):
            if col.get("fieldtype") in ("Currency", "Float", "Int", "Percent"):
                for column_cells in sheet.iter_cols(min_col=idx, max_col=idx, min_row=2):
                    for cell in column_cells:
                        if type(cell.value) in (int, float):
                            cell.number_format = "#,##0.00"
                            cell.alignment = Alignment(horizontal="right")
        sheet.column_dimensions["A"].width = 52
        for idx in range(2, len(columns) + 1):
            sheet.column_dimensions[sheet.cell(row=1, column=idx).column_letter].width = 18
    workbook.save(path)
    return True


# --------------------------------------------------------------------------
# Estado de Cambios en el Patrimonio (matriz por cuenta) - NIC 1
# --------------------------------------------------------------------------


def _balances(company, accounts, date_from=None, date_to=None):
    """Saldo (crédito - débito) por cuenta. Positivo = saldo acreedor."""
    if not accounts:
        return {}
    where = "company=%(company)s and is_cancelled=0 and account in %(accounts)s"
    params = {"company": company, "accounts": tuple(accounts)}
    if date_from and date_to:
        where += " and posting_date between %(from)s and %(to)s"
        params["from"] = d2s(date_from)
        params["to"] = d2s(date_to)
    elif date_to:
        where += " and posting_date <= %(to)s"
        params["to"] = d2s(date_to)
    rows = frappe.db.sql(
        f"""select account, sum(debit) as debit, sum(credit) as credit
            from `tabGL Entry` where {where} group by account""",
        params,
        as_dict=True,
    )
    return {r.account: flt(r.credit) - flt(r.debit) for r in rows}


def _root_balances(company, as_of):
    """Saldo (débito - crédito) acumulado por naturaleza (root_type)."""
    rows = frappe.db.sql(
        """select root_type, sum(debit) - sum(credit) as bal
           from `tabGL Entry`
           where company=%s and is_cancelled=0 and posting_date <= %s
           group by root_type""",
        (company, d2s(as_of)),
        as_dict=True,
    )
    return {r.root_type: flt(r.bal) for r in rows}


def accounting_check(company, as_of):
    """Cuadre contable real: A + G = P + PN + I  (en términos débito-crédito).

    Suma algebraica de todos los movimientos por naturaleza. Si las partidas
    están correctamente clasificadas y el libro cuadra, el resultado es 0.
    """
    bal = _root_balances(company, as_of)
    return sum(bal.values())


def equity_statement(company, from_date, to_date):
    """Devuelve (columns, rows, resultado) con el ECPN matricial (NIC 1)."""
    accounts = frappe.db.sql(
        """select name, account_name, account_number from `tabAccount`
           where company=%s and root_type='Equity' and is_group=0 and disabled=0
           order by ifnull(account_number,''), account_name""",
        company,
        as_dict=True,
    )
    pl_accounts = frappe.db.sql_list(
        """select name from `tabAccount`
           where company=%s and root_type in ('Income','Expense') and is_group=0 and disabled=0""",
        company,
    )

    names = [a.name for a in accounts]
    opening_date = add_days(d2s(from_date), -1)
    opening = _balances(company, names, date_to=opening_date)
    movement = _balances(company, names, date_from=from_date, date_to=to_date)
    pl_movement = _balances(company, pl_accounts, date_from=from_date, date_to=to_date)

    resultado = sum(pl_movement.values())

    columns = [{"fieldname": "concepto", "label": "Concepto", "fieldtype": "Data", "width": 320}]
    for acc in accounts:
        columns.append(
            {
                "fieldname": f"acc_{acc.name}",
                "label": (f"{acc.account_number} " if acc.account_number else "")
                + (acc.account_name or acc.name),
                "fieldtype": "Currency",
                "width": 160,
            }
        )
    columns.append(
        {"fieldname": "resultado", "label": "Resultado del Ejercicio", "fieldtype": "Currency", "width": 170}
    )
    columns.append(
        {"fieldname": "total", "label": "TOTAL PATRIMONIO", "fieldtype": "Currency", "width": 180}
    )

    def total_of(source):
        return sum(source.get(a.name, 0.0) for a in accounts)

    rows = []
    # Saldo inicial
    row = {"concepto": "Saldo inicial del patrimonio", "resultado": 0.0}
    for acc in accounts:
        row[f"acc_{acc.name}"] = opening.get(acc.name, 0.0)
    row["total"] = total_of(opening)
    rows.append(row)

    # Movimientos del período
    row = {"concepto": "Movimientos del período (aportes, reservas, distribuciones)", "resultado": resultado}
    for acc in accounts:
        row[f"acc_{acc.name}"] = movement.get(acc.name, 0.0)
    row["total"] = total_of(movement) + resultado
    rows.append(row)

    # Saldo final
    row = {"concepto": "SALDO FINAL DEL PATRIMONIO NETO", "resultado": resultado}
    for acc in accounts:
        row[f"acc_{acc.name}"] = opening.get(acc.name, 0.0) + movement.get(acc.name, 0.0)
    row["total"] = total_of(opening) + total_of(movement) + resultado
    rows.append(row)

    return columns, rows, resultado


# --------------------------------------------------------------------------
# Libro Diario y Libro Mayor (formato Guatemala)
# --------------------------------------------------------------------------
# En Guatemala el Libro Diario es el registro cronológico de las partidas
# (asientos) con sus cuentas al Debe y al Haber, numeradas por folio y con la
# leyenda "Sumas iguales" en cada partida. El Libro Mayor es el traslado de esas
# partidas a cada cuenta, con su saldo corriente.

LD_COLUMNS = [
    {"fieldname": "folio", "label": "Folio", "fieldtype": "Int", "width": 70},
    {"fieldname": "fecha", "label": "Fecha", "fieldtype": "Data", "width": 100},
    {"fieldname": "documento", "label": "Documento", "fieldtype": "Data", "width": 210},
    {"fieldname": "codigo", "label": "Código", "fieldtype": "Data", "width": 110},
    {"fieldname": "cuenta", "label": "Cuenta", "fieldtype": "Data", "width": 260},
    {"fieldname": "descripcion", "label": "Descripción", "fieldtype": "Data", "width": 320},
    {"fieldname": "debe", "label": "Debe", "fieldtype": "Currency", "width": 130},
    {"fieldname": "haber", "label": "Haber", "fieldtype": "Currency", "width": 130},
]

LM_COLUMNS = [
    {"fieldname": "codigo", "label": "Código", "fieldtype": "Data", "width": 110},
    {"fieldname": "cuenta", "label": "Cuenta", "fieldtype": "Data", "width": 240},
    {"fieldname": "fecha", "label": "Fecha", "fieldtype": "Data", "width": 100},
    {"fieldname": "documento", "label": "Documento", "fieldtype": "Data", "width": 200},
    {"fieldname": "descripcion", "label": "Descripción", "fieldtype": "Data", "width": 300},
    {"fieldname": "debe", "label": "Debe", "fieldtype": "Currency", "width": 130},
    {"fieldname": "haber", "label": "Haber", "fieldtype": "Currency", "width": 130},
    {"fieldname": "saldo", "label": "Saldo", "fieldtype": "Currency", "width": 130},
]

def company_identity(company):
    info = frappe.db.get_value(
        "Company", company, ["company_name", "tax_id", "default_currency"], as_dict=True
    ) or frappe._dict()
    return {
        "name": info.get("company_name") or company,
        "tax_id": info.get("tax_id") or "",
        "currency": info.get("default_currency") or "",
    }


def _gl_rows(company, from_date, to_date):
    return frappe.db.sql(
        """
        select
            g.posting_date,
            g.voucher_type,
            g.voucher_no,
            g.account,
            ifnull(a.account_number, '') as account_number,
            ifnull(a.account_name, g.account) as account_name,
            ifnull(a.root_type, '') as root_type,
            ifnull(g.remarks, '') as remarks,
            ifnull(g.is_opening, 'No') as is_opening,
            g.debit,
            g.credit
        from `tabGL Entry` g
        left join `tabAccount` a on a.name = g.account
        where g.company = %(company)s
          and ifnull(g.is_cancelled, 0) = 0
          and g.posting_date between %(from_date)s and %(to_date)s
        order by g.posting_date, g.creation, g.name
        """,
        {"company": company, "from_date": d2s(from_date), "to_date": d2s(to_date)},
        as_dict=True,
    )


# -------------------------------- Libro Diario ----------------------------


def build_libro_diario(company, from_date, to_date, folio_inicio=1):
    """Agrupa los movimientos en partidas (asientos) cronológicas y las numera.

    Se recorre el libro mayor en orden cronológico y se agrupa por documento de
    origen. Si por alguna razón las líneas de un mismo documento no quedaran
    juntas, el diccionario las vuelve a unir sin perder el orden de aparición.
    """
    asientos = {}
    for row in _gl_rows(company, from_date, to_date):
        key = (row.posting_date, row.voucher_type, row.voucher_no)
        asiento = asientos.get(key)
        if asiento is None:
            asiento = asientos[key] = {
                "key": key,
                "fecha": row.posting_date,
                "voucher_type": row.voucher_type,
                "voucher_no": row.voucher_no,
                "lines": [],
                "debe": 0.0,
                "haber": 0.0,
            }
        asiento["lines"].append(row)
        asiento["debe"] += flt(row.debit)
        asiento["haber"] += flt(row.credit)

    # Los dict conservan el orden de inserción: las partidas quedan en el orden
    # cronológico en que se registraron (igual que el reporte de ERPNext).
    ordered = list(asientos.values())
    for i, asiento in enumerate(ordered, start=folio_inicio):
        asiento["folio"] = i
    return ordered


def libro_diario_flat(asientos):
    out = []
    for asiento in asientos:
        doc = f"{asiento['voucher_type']} {asiento['voucher_no']}".strip()
        fecha = getdate(asiento["fecha"]).strftime("%d/%m/%Y")
        for row in asiento["lines"]:
            out.append(
                {
                    "folio": asiento["folio"],
                    "fecha": fecha,
                    "documento": doc,
                    "codigo": row.account_number,
                    "cuenta": row.account_name,
                    "descripcion": row.remarks,
                    "debe": flt(row.debit),
                    "haber": flt(row.credit),
                }
            )
    return out


def write_libro_diario_html(path, company, from_date, to_date, asientos):
    ident = company_identity(company)
    period = (
        f"Del {getdate(from_date).strftime('%d/%m/%Y')} al {getdate(to_date).strftime('%d/%m/%Y')}"
        f" · {len(asientos)} partidas · moneda {ident['currency']}"
    )

    blocks = []
    total_debe = total_haber = 0.0
    for asiento in asientos:
        lines = "".join(
            "<tr>"
            f'<td class="cod">{_esc(r.account_number)}</td>'
            f"<td>{_esc(r.account_name)}</td>"
            f"<td>{_esc(r.remarks)}</td>"
            f'<td class="num">{_money(r.debit)}</td>'
            f'<td class="num">{_money(r.credit)}</td>'
            "</tr>"
            for r in asiento["lines"]
        )
        total_debe += flt(asiento["debe"])
        total_haber += flt(asiento["haber"])
        blocks.append(
            '<div class="asiento">'
            '<div class="ahead">'
            f'<span class="folio">Folio {asiento["folio"]}</span>'
            f'<span>Asiento No. {asiento["folio"]}</span>'
            f'<span>Fecha: {getdate(asiento["fecha"]).strftime("%d/%m/%Y")}</span>'
            f'<span class="doc">Documento: {_esc(asiento["voucher_type"])} {_esc(asiento["voucher_no"])}</span>'
            "</div>"
            "<table>"
            "<thead><tr><th>Código</th><th>Cuenta</th><th>Descripción</th>"
            '<th class="num">Debe</th><th class="num">Haber</th></tr></thead>'
            f"<tbody>{lines}</tbody>"
            '<tfoot><tr><td colspan="3">Sumas iguales</td>'
            f'<td class="num">{_money(asiento["debe"])}</td>'
            f'<td class="num">{_money(asiento["haber"])}</td></tr></tfoot>'
            "</table></div>"
        )

    diferencia = total_debe - total_haber
    cuadra = abs(diferencia) < 0.01
    resumen = (
        '<div class="panel"><h3>Totales del período</h3>'
        "<table><tbody>"
        f'<tr><td>Total del Debe</td><td class="num">{_money(total_debe)}</td></tr>'
        f'<tr><td>Total del Haber</td><td class="num">{_money(total_haber)}</td></tr>'
        f'<tr class="total"><td>Diferencia (Debe − Haber)</td>'
        f'<td class="num">{_money(diferencia)}</td></tr>'
        "</tbody></table>"
        f'<div class="body"><p>Estado del libro: {_badge(cuadra)}</p></div>'
        "</div>"
    )

    cards = [
        ("Partidas", str(len(asientos))),
        ("Total Debe", total_debe),
        ("Total Haber", total_haber),
        ("Diferencia", diferencia),
    ]

    body = (
        _masthead("Libro Diario", period)
        + '<div class="content">'
        + _cards_html(cards)
        + "".join(blocks)
        + resumen
        + '<div class="note">Generado automáticamente por scripts/estados_financieros.py · '
        f"Diferencia Debe − Haber: {_money(diferencia)}</div>"
        + "</div>"
    )

    with open(path, "w", encoding="utf-8") as handle:
        handle.write(_page("Libro Diario", body))


# -------------------------------- Libro Mayor -----------------------------


def _opening_by_account(company, from_date):
    rows = frappe.db.sql(
        """select g.account, sum(g.debit) - sum(g.credit) as bal
           from `tabGL Entry` g
           where g.company = %(company)s and ifnull(g.is_cancelled, 0) = 0
             and g.posting_date < %(from_date)s
           group by g.account""",
        {"company": company, "from_date": d2s(from_date)},
        as_dict=True,
    )
    return {r.account: flt(r.bal) for r in rows}


def build_libro_mayor(company, from_date, to_date):
    """Traslada las partidas a cada cuenta, con saldo inicial, movimientos y saldo final."""
    opening = _opening_by_account(company, from_date)
    accounts = {}
    for row in _gl_rows(company, from_date, to_date):
        acc = accounts.get(row.account)
        if acc is None:
            acc = accounts[row.account] = {
                "account": row.account,
                "account_number": row.account_number,
                "account_name": row.account_name,
                "root_type": row.root_type,
                "lines": [],
                "debe": 0.0,
                "haber": 0.0,
            }
        acc["lines"].append(row)
        acc["debe"] += flt(row.debit)
        acc["haber"] += flt(row.credit)

    for account, saldo in opening.items():
        if account in accounts or abs(flt(saldo)) < 0.005:
            continue
        info = frappe.db.get_value(
            "Account", account, ["account_number", "account_name", "root_type"], as_dict=True
        ) or frappe._dict()
        accounts[account] = {
            "account": account,
            "account_number": info.get("account_number") or "",
            "account_name": info.get("account_name") or account,
            "root_type": info.get("root_type") or "",
            "lines": [],
            "debe": 0.0,
            "haber": 0.0,
        }

    result = []
    for account in sorted(
        accounts, key=lambda a: (accounts[a]["account_number"] or "", accounts[a]["account_name"] or "")
    ):
        acc = accounts[account]
        running = acc["saldo_inicial"] = flt(opening.get(account, 0.0))
        for row in acc["lines"]:
            running += flt(row.debit) - flt(row.credit)
            row["saldo"] = running
        acc["saldo_final"] = running
        result.append(acc)
    return result


def libro_mayor_flat(accounts):
    out = []
    for acc in accounts:
        base = {
            "codigo": acc["account_number"],
            "cuenta": acc["account_name"],
        }
        out.append({**base, "fecha": "", "documento": "", "descripcion": "Saldo inicial",
                    "debe": "", "haber": "", "saldo": acc["saldo_inicial"]})
        for row in acc["lines"]:
            out.append(
                {
                    **base,
                    "fecha": getdate(row.posting_date).strftime("%d/%m/%Y"),
                    "documento": f"{row.voucher_type} {row.voucher_no}".strip(),
                    "descripcion": row.remarks,
                    "debe": flt(row.debit),
                    "haber": flt(row.credit),
                    "saldo": row.saldo,
                }
            )
        out.append({**base, "fecha": "", "documento": "", "descripcion": "Totales del período",
                    "debe": acc["debe"], "haber": acc["haber"], "saldo": acc["saldo_final"]})
    return out


def write_libro_mayor_html(path, company, from_date, to_date, accounts):
    ident = company_identity(company)
    period = (
        f"Del {getdate(from_date).strftime('%d/%m/%Y')} al {getdate(to_date).strftime('%d/%m/%Y')}"
        f" · {len(accounts)} cuentas · moneda {ident['currency']}"
    )

    tot_ini = sum(flt(acc["saldo_inicial"]) for acc in accounts)
    tot_debe = sum(flt(acc["debe"]) for acc in accounts)
    tot_haber = sum(flt(acc["haber"]) for acc in accounts)
    tot_fin = sum(flt(acc["saldo_final"]) for acc in accounts)
    cuadra = abs(tot_debe - tot_haber) < 0.01 and abs(tot_fin) < 0.01

    # ---- Balance de saldos (una fila por cuenta, con su saldo final) ----
    saldos = "".join(
        "<tr>"
        f'<td class="cod">{_esc(acc["account_number"])}</td>'
        f"<td>{_esc(acc['account_name'])}</td>"
        + _num_cell(acc["saldo_inicial"])
        + _num_cell(acc["debe"])
        + _num_cell(acc["haber"])
        + _num_cell(acc["saldo_final"])
        + "</tr>"
        for acc in accounts
    )
    balance = (
        '<div class="panel"><h3>Balance de saldos</h3>'
        "<table><thead><tr><th>Código</th><th>Cuenta</th>"
        '<th class="num">Saldo inicial</th><th class="num">Debe</th>'
        '<th class="num">Haber</th><th class="num">Saldo final</th></tr></thead>'
        f"<tbody>{saldos}"
        '<tr class="total"><td colspan="2">Totales</td>'
        + _num_cell(tot_ini)
        + _num_cell(tot_debe)
        + _num_cell(tot_haber)
        + _num_cell(tot_fin)
        + "</tr>"
        "</tbody></table>"
        f'<div class="body"><p>Estado del libro: {_badge(cuadra)} '
        '<span class="muted" style="font-size:12px">'
        "(el Debe debe igualar al Haber y la suma de los saldos finales debe dar 0.00)"
        "</span></p></div>"
        "</div>"
    )

    # ---- Detalle por cuenta (saldo corriente y saldo final) ----
    blocks = []
    for acc in accounts:
        rows = [
            "<tr>"
            '<td colspan="3" class="cod">Saldo inicial</td>'
            '<td class="num"></td><td class="num"></td>'
            + _num_cell(acc["saldo_inicial"])
            + "</tr>"
        ]
        for row in acc["lines"]:
            rows.append(
                "<tr>"
                f'<td>{getdate(row.posting_date).strftime("%d/%m/%Y")}</td>'
                f'<td>{_esc(row.voucher_type)} {_esc(row.voucher_no)}</td>'
                f"<td>{_esc(row.remarks)}</td>"
                + _num_cell(row.debit)
                + _num_cell(row.credit)
                + _num_cell(row.saldo)
                + "</tr>"
            )
        blocks.append(
            '<div class="acc">'
            f'<h2><span class="code">{_esc(acc["account_number"])}</span>{_esc(acc["account_name"])}</h2>'
            "<table><thead><tr><th>Fecha</th><th>Documento</th><th>Descripción</th>"
            '<th class="num">Debe</th><th class="num">Haber</th><th class="num">Saldo</th></tr></thead>'
            f'<tbody>{"".join(rows)}</tbody>'
            '<tfoot><tr><td colspan="3">Totales del período</td>'
            + _num_cell(acc["debe"])
            + _num_cell(acc["haber"])
            + _num_cell(acc["saldo_final"])
            + "</tr></tfoot>"
            "</table></div>"
        )

    cards = [
        ("Cuentas", str(len(accounts))),
        ("Total Debe", tot_debe),
        ("Total Haber", tot_haber),
        ("Saldo final (Σ)", tot_fin),
    ]

    body = (
        _masthead("Libro Mayor", period)
        + '<div class="content">'
        + _cards_html(cards)
        + balance
        + '<h2 class="sec">Detalle por cuenta</h2>'
        + "".join(blocks)
        + '<div class="note">Generado automáticamente por scripts/estados_financieros.py · '
        f'Diferencia Debe − Haber: {_money(tot_debe - tot_haber)} · '
        f"Suma de saldos finales: {_money(tot_fin)}</div>"
        + "</div>"
    )

    with open(path, "w", encoding="utf-8") as handle:
        handle.write(_page("Libro Mayor", body))


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------


def _emit(slug, title, columns, rows, note="", cards=None):
    if not columns:
        print(f"  ! {title}: sin datos/columnas")
        return
    write_csv(os.path.join(OUTDIR, f"{slug}.csv"), columns, rows)
    write_html(os.path.join(OUTDIR, f"{slug}.html"), title, SUBTITLE, columns, rows, note, cards)
    print(f"  ✔ {title}  ({len(rows)} filas)")


def _safe_report(report, filters, slug, title, note="", cards=None):
    """Ejecuta un reporte; si falla, avisa pero no interrumpe el resto.

    'cards' puede ser una lista de (etiqueta, valor) o una función que recibe
    las filas y devuelve esa lista (para las tarjetas de resumen del encabezado).
    """
    try:
        cols, rows = run_report(report, filters)
    except Exception as exc:  # noqa: BLE001
        print(f"  ! {title}: no se pudo ejecutar '{report}' ({exc})")
        return [], []
    resolved = cards(rows) if callable(cards) else cards
    _emit(slug, title, cols, rows, note, resolved)
    return cols, rows


def main():
    global SUBTITLE, IDENTITY

    company = resolve_company()
    fy_doc = resolve_fiscal_year(company, os.environ.get("EF_FISCAL_YEAR"))
    fy = fy_doc.name
    from_date = fy_doc.year_start_date
    to_date = fy_doc.year_end_date

    currency = frappe.get_cached_value("Company", company, "default_currency")
    IDENTITY = company_identity(company)
    SUBTITLE = (
        f"{company} · ejercicio fiscal {fy} · {getdate(from_date)} a {getdate(to_date)}"
        f" · moneda {currency}"
    )
    note = "Generado automáticamente por scripts/estados_financieros.py"

    os.makedirs(OUTDIR, exist_ok=True)
    sheets = []
    summary = []

    print("=" * 90)
    print(" ESTADOS FINANCIEROS COMPLETOS")
    print(f" {SUBTITLE}")
    print("=" * 90)

    # ---- 1. Balance General -------------------------------------------------
    cols, rows = _safe_report(
        "Balance Sheet",
        fin_filters(company, from_date, to_date, fy),
        "01-balance-general",
        "Balance General (Estado de Situación Financiera)",
        note,
        cards=lambda rows: [
            ("Total activo", _pick_total(rows, "Total Assets", "Total Activo")),
            ("Total pasivo", _pick_total(rows, "Total Liabilities", "Total Pasivo")),
            ("Total patrimonio", _pick_total(rows, "Total Equity", "Total Patrimonio")),
        ],
    )
    if rows:
        sheets.append(("Balance General", cols, rows))
        summary.append(("Total activo", _pick_total(rows, "Total Assets", "Total Activo")))
        summary.append(("Total pasivo", _pick_total(rows, "Total Liabilities", "Total Pasivo")))
        summary.append(("Total patrimonio", _pick_total(rows, "Total Equity", "Total Patrimonio")))
        summary.append(
            (
                "Balance General · Activo − (Pasivo + Patrimonio)",
                flt(_pick_total(rows, "Total Assets", "Total Activo"))
                - flt(_pick_total(rows, "Total Liabilities", "Total Pasivo"))
                - flt(_pick_total(rows, "Total Equity", "Total Patrimonio")),
            )
        )

    # ---- 2. Estado de Resultados -------------------------------------------
    cols, rows = _safe_report(
        "Profit and Loss Statement",
        fin_filters(company, from_date, to_date, fy, accumulated_values=0),
        "02-estado-resultados",
        "Estado de Resultados",
        note,
        cards=lambda rows: [
            ("Total ingresos", _pick_total(rows, "Total Income", "Total Ingresos")),
            ("Total gastos", _pick_total(rows, "Total Expense", "Total Gastos")),
            (
                "Utilidad / (Pérdida)",
                _pick_total(rows, "Profit for the year", "Net Profit", "Total", "Utilidad"),
            ),
        ],
    )
    if rows:
        sheets.append(("Estado de Resultados", cols, rows))
        summary.append(
            (
                "Utilidad / (Pérdida) neta",
                _pick_total(rows, "Profit for the year", "Net Profit", "Total", "Utilidad"),
            )
        )

    # ---- 3. Balance Financiero (Balanza de Comprobación) -------------------
    tb_filters = frappe._dict(
        {
            "company": company,
            "fiscal_year": fy,
            "from_date": d2s(from_date),
            "to_date": d2s(to_date),
            "period_start_date": d2s(from_date),
            "period_end_date": d2s(to_date),
            "filter_based_on": "Fiscal Year",
            "periodicity": "Yearly",
            "with_period_closing_entry_for_opening": 1,
            "with_period_closing_entry_for_current_period": 1,
            "show_zero_values": 0,
            "show_unclosed_fy_pl_balances": 1,
            "include_default_book_entries": 1,
            "show_net_values": 1,
            "show_group_accounts": 1,
        }
    )
    cols, rows = _safe_report(
        "Trial Balance",
        tb_filters,
        "03-balance-financiero",
        "Balance Financiero (Balanza de Comprobación)",
        note,
    )
    if rows:
        sheets.append(("Balanza de Comprobacion", cols, rows))

    # ---- 4. Estado de Flujo de Efectivo ------------------------------------
    cols, rows = _safe_report(
        "Cash Flow",
        fin_filters(
            company,
            from_date,
            to_date,
            fy,
            show_opening_and_closing_balance=1,
            show_account_details="Summary",
        ),
        "04-flujo-efectivo",
        "Estado de Flujo de Efectivo",
        note,
        cards=lambda rows: [
            ("Flujo neto del período", _pick_total(rows, "Net Cash", "Net Cash Flow", "Net Change in Cash")),
        ],
    )
    if rows:
        sheets.append(("Flujo de Efectivo", cols, rows))
        summary.append(
            (
                "Flujo neto del período",
                _pick_total(rows, "Net Cash", "Net Cash Flow", "Net Change in Cash"),
            )
        )

    # ---- 5. Estado de Cambios en el Patrimonio (matriz) --------------------
    try:
        cols, rows, resultado = equity_statement(company, from_date, to_date)
        control = accounting_check(company, to_date)
        _emit(
            "05-estado-patrimonio",
            "Estado de Cambios en el Patrimonio Neto (NIC 1)",
            cols,
            rows,
            note=f"{note} · Cuadre contable del libro (debe ser 0): {flt(control, 2):,.2f}",
            cards=[
                ("Resultado del ejercicio", resultado),
                ("Cuadre contable del libro (debe ser 0)", control),
            ],
        )
        sheets.append(("Patrimonio Neto", cols, rows))
        summary.append(("Resultado del ejercicio (según mayor)", resultado))
        summary.append(("Cuadre contable del libro (debe ser 0)", control))
    except Exception as exc:  # noqa: BLE001
        print(f"  ! No se pudo construir el ECPN matricial: {exc}")

    # ---- 5b. Plantilla nativa ERPNext (Custom Financial Statement) ---------
    template = "Estado de Cambios en el Patrimonio Neto (NIC 1)"
    if frappe.db.exists("Financial Report Template", template):
        _safe_report(
            "Custom Financial Statement",
            fin_filters(company, from_date, to_date, fy, report_template=template),
            "05b-estado-patrimonio-plantilla",
            "Estado de Cambios en el Patrimonio Neto (plantilla ERPNext)",
            note,
        )
    else:
        print("  ! La plantilla ECPN no existe (ejecuta setup_accounting.py)")

    # ---- 6. Libro Diario (Guatemala) ---------------------------------------
    try:
        folio_inicio = int(os.environ.get("LD_FOLIO_INICIO") or 1)
    except ValueError:
        folio_inicio = 1
    try:
        asientos = build_libro_diario(company, from_date, to_date, folio_inicio)
        ld_rows = libro_diario_flat(asientos)
        write_libro_diario_html(
            os.path.join(OUTDIR, "06-libro-diario.html"), company, from_date, to_date, asientos
        )
        write_csv(os.path.join(OUTDIR, "06-libro-diario.csv"), LD_COLUMNS, ld_rows)
        print(f"  ✔ Libro Diario  ({len(asientos)} partidas, {len(ld_rows)} líneas)")
        if ld_rows:
            sheets.append(("Libro Diario", LD_COLUMNS, ld_rows))
        ld_debe = sum(flt(a["debe"]) for a in asientos)
        ld_haber = sum(flt(a["haber"]) for a in asientos)
        summary.append(("Libro Diario · total Debe", ld_debe))
        summary.append(("Libro Diario · total Haber", ld_haber))
        summary.append(("Libro Diario · diferencia (debe ser 0)", ld_debe - ld_haber))
    except Exception as exc:  # noqa: BLE001
        print(f"  ! No se pudo construir el Libro Diario: {exc}")

    # ---- 7. Libro Mayor (Guatemala) ----------------------------------------
    try:
        mayor = build_libro_mayor(company, from_date, to_date)
        lm_rows = libro_mayor_flat(mayor)
        write_libro_mayor_html(
            os.path.join(OUTDIR, "07-libro-mayor.html"), company, from_date, to_date, mayor
        )
        write_csv(os.path.join(OUTDIR, "07-libro-mayor.csv"), LM_COLUMNS, lm_rows)
        print(f"  ✔ Libro Mayor  ({len(mayor)} cuentas, {len(lm_rows)} líneas)")
        if lm_rows:
            sheets.append(("Libro Mayor", LM_COLUMNS, lm_rows))
        lm_debe = sum(flt(a["debe"]) for a in mayor)
        lm_haber = sum(flt(a["haber"]) for a in mayor)
        lm_saldo = sum(flt(a["saldo_final"]) for a in mayor)
        summary.append(("Libro Mayor · total Debe", lm_debe))
        summary.append(("Libro Mayor · total Haber", lm_haber))
        summary.append(("Libro Mayor · suma de saldos finales (debe ser 0)", lm_saldo))
    except Exception as exc:  # noqa: BLE001
        print(f"  ! No se pudo construir el Libro Mayor: {exc}")

    # ---- índice + xlsx ------------------------------------------------------
    _write_index(SUBTITLE)

    if summary:
        sheets.append(
            (
                "Resumen y cuadre",
                [
                    {"fieldname": "concepto", "label": "Concepto", "fieldtype": "Data", "width": 460},
                    {"fieldname": "valor", "label": "Importe", "fieldtype": "Currency", "width": 170},
                ],
                [{"concepto": key, "valor": flt(value)} for key, value in summary if value is not None],
            )
        )
    if sheets and write_xlsx(os.path.join(OUTDIR, "estados-financieros.xlsx"), sheets):
        print("  ✔ estados-financieros.xlsx")

    print("-" * 90)
    for label, value in summary:
        if value is None:
            continue
        print(f" {label[:58]:<58} {flt(value, 2):>18,.2f}")
    print("-" * 90)
    print(f" Archivos en: {OUTDIR}")
    print(f"{SENTINEL} {company} {fy}")
    return 0


def _pick_total(rows, *labels):
    """Busca una fila de total por su etiqueta y devuelve su importe."""
    meta_keys = {
        "indent",
        "is_group",
        "is_total_row",
        "has_value",
        "warn_if_negative",
        "account_name",
    }

    def norm(text):
        return "".join(ch for ch in cstr(text).lower() if ch.isalnum())

    wanted = [norm(label) for label in labels]
    for row in rows or []:
        label = ""
        for key in ("account", "acc_name", "account_name", "label"):
            if row.get(key):
                label = cstr(row.get(key))
                break
        name = norm(label)
        if not name or not any(w == name or w in name for w in wanted):
            continue
        for key in ("total", "closing_credit", "closing_debit", "balance"):
            value = row.get(key)
            if type(value) in (int, float):
                return value
        numbers = [
            v
            for k, v in row.items()
            if k not in meta_keys and type(v) in (int, float)
        ]
        if numbers:
            return numbers[-1]
    return None


def _write_index(subtitle):
    entries = [
        ("01-balance-general.html", "Balance General"),
        ("02-estado-resultados.html", "Estado de Resultados"),
        ("03-balance-financiero.html", "Balance Financiero (Balanza)"),
        ("04-flujo-efectivo.html", "Estado de Flujo de Efectivo"),
        ("05-estado-patrimonio.html", "Estado de Cambios en el Patrimonio Neto"),
        ("05b-estado-patrimonio-plantilla.html", "ECPN (plantilla ERPNext)"),
        ("06-libro-diario.html", "Libro Diario (Guatemala)"),
        ("07-libro-mayor.html", "Libro Mayor (Guatemala)"),
    ]
    files = [(h, t) for h, t in entries if os.path.exists(os.path.join(OUTDIR, h))]
    libros = [(h, t) for h, t in files if h[:2] in ("06", "07")]
    estados = [(h, t) for h, t in files if h[:2] not in ("06", "07")]

    def items(pairs):
        return "".join(
            f'<li><a href="{html.escape(h)}">{html.escape(t)}</a> '
            f'<span class="muted">· <a href="{html.escape(h.replace(".html", ".csv"))}">csv</a></span></li>'
            for h, t in pairs
        )

    body = (
        _masthead("Estados financieros y libros contables", subtitle)
        + '<div class="content">'
        + '<h2 class="sec">Estados financieros</h2>'
        + f"<ul>{items(estados)}</ul>"
        + '<h2 class="sec">Libros contables (Guatemala)</h2>'
        + f"<ul>{items(libros)}</ul>"
        + '<div class="note"><a href="estados-financieros.xlsx">Descargar todo en Excel '
        "(estados-financieros.xlsx)</a> · Imprime cualquier reporte con Ctrl+P (→ PDF).</div>"
        + "</div>"
    )
    with open(os.path.join(OUTDIR, "index.html"), "w", encoding="utf-8") as handle:
        handle.write(_page("Estados financieros", body))


if __name__ == "__main__":
    try:
        main()
    except Exception as _exc:  # noqa: BLE001
        import traceback

        traceback.print_exc()
        print(f"ERROR: {_exc}")
        raise SystemExit(1)
